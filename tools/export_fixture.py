"""Export the restaurant tab from a downloaded .xlsx into tests/fixture.json,
shaped like the data Apps Script hands to parseSheet(). Used only for local
testing; fixture.json is gitignored because it holds the real list.

Usage:  python tools/export_fixture.py "path/to/Seattle vibes _ restaurants .xlsx"
"""
import json
import sys
from pathlib import Path

import openpyxl

TAB = "Seattle dining"
COLS = 5
OUT = Path(__file__).resolve().parent.parent / "tests" / "fixture.json"


def main(xlsx_path: str) -> None:
    ws = openpyxl.load_workbook(xlsx_path)[TAB]
    data = {"values": [], "backgrounds": [], "bold": [], "strike": [], "links": []}
    for r in range(1, ws.max_row + 1):
        cells = [ws.cell(r, c) for c in range(1, COLS + 1)]
        name = cells[0]
        data["values"].append(["" if c.value is None else str(c.value) for c in cells])
        fill = name.fill.fgColor.rgb if name.fill and name.fill.fill_type else None
        data["backgrounds"].append("#" + fill[-6:].lower() if isinstance(fill, str) else "#ffffff")
        data["bold"].append(bool(name.font.bold))
        data["strike"].append(bool(name.font.strike))
        data["links"].append(name.hyperlink.target if name.hyperlink and name.hyperlink.target else "")
    OUT.parent.mkdir(exist_ok=True)
    OUT.write_text(json.dumps(data), encoding="utf-8")
    print(f"Wrote {len(data['values'])} rows to {OUT}")


if __name__ == "__main__":
    main(sys.argv[1])
